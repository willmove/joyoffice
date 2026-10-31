"""Native M9 canvas and M10 clock acceptance, plus independent public PPTX/XLSX readers."""
import argparse,hashlib,importlib.util,json,os,subprocess,time,zipfile,io
from pathlib import Path
from paths import ROOT, SCRIPTS, FIXTURES, EVIDENCE, LICENSES, binary, native_helper
spec=importlib.util.spec_from_file_location('bench',SCRIPTS / 'benchmark-m4.py');bench=importlib.util.module_from_spec(spec);spec.loader.exec_module(bench)
def main():
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('--binary',type=Path,default=binary('joydeck', 'release'));p.add_argument('--output',type=Path,default=EVIDENCE / 'm10'/f'run-{time.time_ns()}');a=p.parse_args();out=a.output.resolve();out.mkdir(parents=True,exist_ok=False);quality={};seen=set()
 with (out/'native.log').open('wb') as log:
  proc=subprocess.Popen([str(a.binary.resolve()),'--authoring-smoke'],cwd=ROOT,env=dict(os.environ,JOYDECK_M9_OUTPUT=str(out)),stdout=log,stderr=subprocess.STDOUT)
  try:
   deadline=time.monotonic()+80
   while time.monotonic()<deadline:
    text=(out/'native.log').read_text(errors='replace', encoding='utf-8')
    if 'failed=' in text or 'panicked' in text:raise RuntimeError(text[-2500:])
    for line in text.splitlines():
     if not line.startswith('M9_CAPTURE '):continue
     _,backend,bounds=line.split(' ',2)
     if backend not in seen:
      quality[backend]=bench.capture(proc.pid,out,f'M9-{backend}',json.loads(bounds));seen.add(backend);(out/f'capture-{0 if backend=="GPUI" else 1}.ack').touch()
    if proc.poll() is not None:
     assert proc.returncode==0 and 'M10_SMOKE clock_pause_blank_navigation_snapshot=true complete=true' in text,text[-2500:];break
    time.sleep(.1)
   else:raise TimeoutError('native authoring validation timeout')
  finally:
   if proc.poll() is None:proc.terminate();proc.wait(timeout=5)
 from pptx import Presentation
 deck=Presentation(out/'authoring-public.pptx');charts=[shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart];assert len(charts)==7
 assert list(charts[0].series[0].values)==[42,58,72,131]
 assert list(charts[4].series[0].values)==[42,58,72,91]
 from pptx.enum.chart import XL_CHART_TYPE
 assert charts[5].chart_type==XL_CHART_TYPE.AREA and charts[6].chart_type==XL_CHART_TYPE.DOUGHNUT
 assert sum(shape.shape_type==6 for slide in deck.slides for shape in slide.shapes)==2
 import xml.etree.ElementTree as E
 z=zipfile.ZipFile(out/'authoring-public.pptx');workbooks=[n for n in z.namelist() if n.endswith('.xlsx')];assert len(workbooks)==7
 chartns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
 rings=[E.fromstring(z.read(n)).find('.//c:doughnutChart/c:holeSize',chartns) for n in z.namelist() if n.startswith('ppt/charts/') and n.endswith('.xml')]
 assert [n.get('val') for n in rings if n is not None]==['50']
 ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
 book=zipfile.ZipFile(io.BytesIO(z.read('ppt/embeddings/chart-1-3.xlsx')));r=E.fromstring(book.read('xl/worksheets/sheet1.xml'));assert r.find('.//s:c[@r="B5"]/s:v',ns).text=='131'
 from pypdf import PdfReader
 assert len(PdfReader(out/'authoring.pdf').pages)==2
 show=json.loads((out/'native-show.json').read_text(encoding='utf-8'));assert show['renderedFrames']>5 and show['snapshotUnchanged'] and show['monotonicPause']
 result={'binary':str(a.binary.resolve()),'sha256':hashlib.sha256(a.binary.read_bytes()).hexdigest(),'canvas':quality,'publicCharts':len(charts),'embeddedWorkbooks':len(workbooks),'cacheWorkbookValue':131,'nativeShow':show,'passed':len(quality)==2 and all(x['passed'] for x in quality.values())}
 (out/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n', encoding='utf-8');print(json.dumps(result,ensure_ascii=False,indent=2));print(out)
 if not result['passed']:raise SystemExit(1)
if __name__=='__main__':main()
