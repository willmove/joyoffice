"""M8 native window/CPU/GPUI validation with independent public DrawingML table readers."""
import argparse, importlib.util, json, os, subprocess, time, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from paths import ROOT, SCRIPTS, FIXTURES, EVIDENCE, LICENSES, binary, native_helper
spec=importlib.util.spec_from_file_location('m4',SCRIPTS / 'benchmark-m4.py');bench=importlib.util.module_from_spec(spec);spec.loader.exec_module(bench)

def public_copy(source,target):
    with zipfile.ZipFile(source) as z: parts={n:z.read(n) for n in z.namelist() if not n.startswith('customXml/')}
    for name in ['_rels/.rels','ppt/_rels/presentation.xml.rels','[Content_Types].xml']:
        root=ET.fromstring(parts[name])
        for child in list(root):
            if 'customXml' in child.attrib.get('Type','') or 'customXml/' in child.attrib.get('Target','') or 'customXml/' in child.attrib.get('PartName',''):root.remove(child)
        parts[name]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--poppler',type=Path,default=Path('/Users/willm/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/bin'));p.add_argument('--binary',type=Path,default=binary('joydeck', 'debug'));p.add_argument('--output',type=Path,default=EVIDENCE / 'm8'/f'run-{time.time_ns()}');a=p.parse_args();out=a.output.resolve();out.mkdir(parents=True,exist_ok=False);quality={};seen=set()
    with (out/'native.log').open('wb') as log:
        process=subprocess.Popen([str(a.binary.resolve()),'--table-smoke'],cwd=ROOT,env=dict(os.environ,JOYDECK_M8_OUTPUT=str(out)),stdout=log,stderr=subprocess.STDOUT)
        try:
            deadline=time.monotonic()+55
            while time.monotonic()<deadline:
                text=(out/'native.log').read_text(errors='replace', encoding='utf-8')
                if 'failed=' in text or 'startup/runtime failure' in text:raise RuntimeError(text[-3000:])
                for line in text.splitlines():
                    if not line.startswith('M8_CAPTURE '):continue
                    _,backend,bounds=line.split(' ',2)
                    if backend in seen:continue
                    quality[backend]=bench.capture(process.pid,out,f'M8-{backend}',json.loads(bounds));seen.add(backend);(out/f'capture-{0 if backend=="GPUI" else 1}.ack').touch()
                if process.poll() is not None:
                    if process.returncode or 'M8_SMOKE complete=true' not in text:raise RuntimeError(text[-3000:])
                    break
                time.sleep(.1)
            else:raise RuntimeError('M8 native validation timed out')
        finally:
            if process.poll() is None:process.terminate();process.wait(timeout=5)
    from pptx import Presentation
    independent={}
    for name in ['table','merged']:
        source=out/f'{name}.pptx';public=out/f'{name}-public.pptx';public_copy(source,public);deck=Presentation(public);tables=[s.table for s in deck.slides[0].shapes if s.has_table];assert len(tables)==1;t=tables[0];assert (len(t.rows),len(t.columns))==(4,5);assert t.cell(0,0).text=='业务板块';assert t.cell(1,2).text=='136'
        merged=name=='merged';assert t.cell(2,3).is_merge_origin==merged
        if merged:assert (t.cell(2,3).span_height,t.cell(2,3).span_width)==(2,2) and t.cell(3,4).is_spanned
        independent[name]={'publicTable':True,'rows':4,'columns':5,'merge':merged,'headerBold':t.cell(0,0).text_frame.paragraphs[0].runs[0].font.bold}
    from pypdf import PdfReader
    assert len(PdfReader(out/'table.pdf').pages)==1
    textbin=a.poppler/'pdftotext'
    if not textbin.exists():textbin=a.poppler.parent/'poppler/bin/pdftotext'
    subprocess.run([str(textbin),'-layout',str(out/'table.pdf'),str(out/'pdf-text.txt')],check=True)
    text=(out/'pdf-text.txt').read_text(encoding='utf-8');assert '业务板块' in text and '企业服务' in text
    result={'binary':str(a.binary.resolve()),'pixelChecks':quality,'publicPptx':independent,'pdfCellText':True,'passed':len(quality)==2 and all(v['passed'] for v in quality.values()),'scope':'macOS ARM64 native adapters and independent public readers; other platforms and physical keyboard unverified'}
    (out/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n', encoding='utf-8');print(json.dumps(result,ensure_ascii=False,indent=2));print(out)
    if not result['passed']:raise SystemExit(1)
if __name__=='__main__':main()
