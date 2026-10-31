"""macOS M7 final-binary fullscreen, public PPTX, image and PDF verification.
Requires Pillow, numpy, python-pptx, pypdf and Poppler.
"""
import argparse,importlib.util,json,os,subprocess,time
from pathlib import Path
from PIL import Image
import numpy as np
from pptx import Presentation
from pypdf import PdfReader
from paths import ROOT, SCRIPTS, FIXTURES, EVIDENCE, LICENSES, binary, native_helper
spec=importlib.util.spec_from_file_location('bench',SCRIPTS / 'benchmark-m4.py');bench=importlib.util.module_from_spec(spec);spec.loader.exec_module(bench)

def capture(pid,title,path):
    windows=[w for w in bench.windows(pid) if w.get('kCGWindowName')==title]
    if len(windows)!=1:raise RuntimeError(f'Expected one {title}: {windows}')
    subprocess.run(['screencapture','-x','-o','-l',str(windows[0]['kCGWindowNumber']),str(path)],check=True)
    return windows[0]

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--binary',type=Path,default=binary('joydeck', 'debug'));p.add_argument('--output',type=Path,default=EVIDENCE / 'm7'/f'run-{time.time_ns()}');p.add_argument('--poppler',type=Path,default=Path('/Users/willm/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/bin'))
    args=p.parse_args();out=args.output.resolve();out.mkdir(parents=True,exist_ok=False)
    with (out/'native.log').open('wb') as log:
        app=subprocess.Popen([str(args.binary.resolve()),'--delivery-smoke'],cwd=ROOT,env=dict(os.environ,JOYDECK_M7_OUTPUT=str(out)),stdout=log,stderr=subprocess.STDOUT)
        try:
            deadline=time.monotonic()+100;show=False;editor=False
            while time.monotonic()<deadline:
                text=(out/'native.log').read_text(errors='replace')
                if 'failed=' in text or 'panicked at' in text or 'startup/runtime failure' in text:raise RuntimeError(text[-3000:])
                if 'M7_SHOW_CAPTURE' in text and not show:
                    # macOS may return the pre-transition backing store after
                    # is_fullscreen() and CPU image readiness are already true.
                    # Wait for the actual captured first frame, with the same
                    # strict pixel threshold and a bounded five-second deadline.
                    showviewport=json.loads(next(line.split(' ',1)[1] for line in text.splitlines() if line.startswith('M7_SHOW_CAPTURE ')))
                    reference=Image.open(out/'png/slide-001.png').convert('RGB')
                    capture_start=time.monotonic();attempts=[]
                    while True:
                        showwindow=capture(app.pid,'JoyDeck — Static Slideshow',out/'slideshow.png')
                        captured=bench.srgb_image(out/'slideshow.png');w,h=captured.size
                        dpi=w/showwindow['kCGWindowBounds']['Width'];vw=showviewport['width']*dpi;vh=showviewport['height']*dpi;title=h-vh
                        targetw=min(vw,vh*960/512);targeth=targetw*512/960
                        box=(round((vw-targetw)/2),round(title+(vh-targeth)/2),round((vw+targetw)/2),round(title+(vh+targeth)/2))
                        crop=captured.crop(box).resize(reference.size,Image.Resampling.BOX)
                        error=float(np.abs(np.asarray(crop,dtype=np.int16)-np.asarray(reference,dtype=np.int16)).mean())
                        attempts.append({'seconds':round(time.monotonic()-capture_start,3),'meanError':error})
                        if error<5.:break
                        if time.monotonic()-capture_start>=5.:raise AssertionError(f'Fullscreen first frame timeout: {attempts}')
                        if len(attempts)==1:captured.save(out/'slideshow-first-capture.png')
                        time.sleep(.15)
                    (out/'capture-readiness.json').write_text(json.dumps(attempts,indent=2)+'\n')
                    (out/'show.ack').touch();show=True
                if 'M7_EDITOR_CAPTURE' in text and not editor:
                    capture(app.pid,'JoyDeck — M7 Acceptance Editor',out/'editor.png');(out/'editor.ack').touch();editor=True
                if app.poll() is not None:
                    if app.returncode!=0 or 'fullscreen_snapshot_hidden_navigation_escape=true complete=true' not in text:raise RuntimeError(text[-3000:])
                    break
                time.sleep(.1)
            else:raise RuntimeError('M7 native timeout')
        finally:
            if app.poll() is None:app.terminate();app.wait(timeout=5)
    # Independent public-format readers do not consult JoyDeck's private snapshot.
    deck=Presentation(out/'slides.pptx');assert len(deck.slides)==3
    assert deck.slides[1]._element.get('show')=='0'
    text='\n'.join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
    assert 'M7 文件交付' in text and '中文导出与静态放映' in text
    assert sum(s.shape_type==13 for s in deck.slides[0].shapes)==1
    assert '演讲备注' in deck.slides[0].notes_slide.notes_text_frame.text
    for file in sorted((out/'corpus').glob('*.pptx')):
        corpus=Presentation(file);assert len(corpus.slides)==1
        assert any('原创语料' in s.text for s in corpus.slides[0].shapes if s.has_text_frame)
    assert len(list((out/'corpus').glob('*.pptx')))==40
    png=Image.open(out/'png/slide-001.png');jpeg=Image.open(out/'jpeg/slide-001.jpg');assert png.size==(1920,1024) and jpeg.size==png.size
    assert abs(png.info['dpi'][0]-144)<.1 and abs(jpeg.info['dpi'][0]-144)<.1
    pdfs={}
    for name in ['slides','notes']:
        pdf=PdfReader(out/f'{name}.pdf');assert len(pdf.pages)==2;assert len(pdf.outline)==2
        textbin=args.poppler/'pdftotext'
        if not textbin.exists():textbin=args.poppler.parent/'poppler/bin/pdftotext'
        subprocess.run([str(textbin),'-layout',str(out/f'{name}.pdf'),str(out/f'{name}.txt')],check=True)
        text=(out/f'{name}.txt').read_text();assert '中文导出与静态放映' in text
        if name=='notes':assert '演讲备注' in text
        subprocess.run([str(args.poppler/'pdftoppm'),'-f','1','-singlefile','-r','144','-png',str(out/f'{name}.pdf'),str(out/f'{name}-render')],check=True)
        pdfs[name]={'pages':2,'outline':2,'unicodeText':True,'sizePoints':list(pdf.pages[0].mediabox)}
    actual=np.asarray(Image.open(out/'slides-render.png').convert('RGB'),dtype=np.int16);reference=np.asarray(png.convert('RGB'),dtype=np.int16);assert actual.shape==reference.shape
    error=float(np.abs(actual-reference).mean());assert error<1.,error
    showimage=bench.srgb_image(out/'slideshow.png');w,h=showimage.size;aspect=960/512
    dpi=w/showwindow['kCGWindowBounds']['Width'];vw=showviewport['width']*dpi;vh=showviewport['height']*dpi;title=h-vh
    targetw=min(vw,vh*aspect);targeth=targetw/aspect;box=(round((vw-targetw)/2),round(title+(vh-targeth)/2),round((vw+targetw)/2),round(title+(vh+targeth)/2))
    crop=showimage.crop(box).resize(png.size,Image.Resampling.BOX);crop.save(out/'slideshow-canvas.png');showerror=float(np.abs(np.asarray(crop,dtype=np.int16)-reference).mean());assert showerror<5.,showerror
    result={'binary':str(args.binary.resolve()),'passed':True,'publicPptxReader':'python-pptx','syntheticCorpusPublicReaderCount':40,'pdf':pdfs,'pdfPngMeanError':error,'slideshowMeanError':showerror,'imageDpi':144,'scope':'macOS adapter/fullscreen; PowerPoint visual and other platforms recorded separately'}
    (out/'verification.json').write_text(json.dumps(result,indent=2,default=float)+'\n');print(json.dumps(result,indent=2,default=float));print(out)
if __name__=='__main__':main()
